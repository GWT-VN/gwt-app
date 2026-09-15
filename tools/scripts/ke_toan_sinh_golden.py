"""Sinh fixture test parity cho engine TypeScript từ tool Python cũ (data/ke-toan, gitignore).
Chạy TAY trên máy có data: python tools/scripts/ke_toan_sinh_golden.py
Che PII: bỏ MST/địa chỉ/người mua; seller là tên công ty (giữ); nếu seller không có chữ
"CÔNG TY|CTY|TNHH|CỔ PHẦN|CP|DNTN|HỘ KINH DOANH" thì thay bằng "NCC-<số>" (cá nhân).
Số 10 chữ số dạng SĐT VN trong tên hàng → 0900000000 (dải giả của scan_pii_secrets.py).

Bổ sung (rà tay 04/09): hoá đơn Be/GSM (grab-like) nhét PII cá nhân bên thứ ba vào `Tên hàng
hóa, dịch vụ` — tên + địa chỉ người gửi hàng, điểm đi/đến, biển số xe. Cắt đuôi từ marker
PII đầu tiên (không phải phần "Tên hàng hóa: <loại>" đứng trước nó). Kiểm chứng: 144 dòng
dính marker này đều nguồn `rule_ncc`/`hoc_ncc` (phân loại theo NCC, không đọc phần đuôi bị
cắt) và `_is_goods()` thoát sớm nhờ từ khoá "cuoc"/"phi " nằm ở phần đầu còn giữ nguyên — nên
cắt đuôi không đổi `expected` đã tính trên desc gốc.
"""
import glob, json, os, re, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PKG = glob.glob(os.path.join(ROOT, "data", "ke-toan", "*", "*", "gwt_ketoan"))[0]
sys.path.insert(0, os.path.dirname(PKG))
from gwt_ketoan import engine, nexia  # noqa: E402

XLSX = glob.glob(os.path.join(ROOT, "data", "ke-toan", "**", "2026", "08.2026 - GWT - NEXIA.xlsx"), recursive=True)[0]
OUT = os.path.join(ROOT, "apps", "web", "lib", "ke-toan", "__fixtures__")
os.makedirs(OUT, exist_ok=True)

CTY = re.compile(r"c[ôo]ng ty|cty|tnhh|c[ổo] ph[ầa]n|\bcp\b|dntn|h[ộo] kinh doanh|bank|ltd|co\.|corp", re.I)
PHONE = re.compile(r"(?<!\d)0[35789]\d{8}(?!\d)")
TAIL_PII = re.compile(
    r"\s*[\(\-;]*\s*(?:BKS|Điểm đi|Điểm đến|Tên người gửi hàng|Tên người gửi|"
    r"Địa chỉ người gửi hàng|Địa chỉ gửi hàng|MST người gửi hàng|"
    r"Mã định danh người gửi hàng)\s*:.*$",
    re.I | re.S,
)


def che(s):
    s = PHONE.sub("0900000000", str(s or ""))
    return TAIL_PII.sub("", s).rstrip(" -;(")


con = engine.connect()
clf = engine.Classifier(con, learn=True)
kmname = {r["ma"]: r["ten"] for r in con.execute("SELECT ma,ten FROM kmcp")}

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["HĐ đầu vào"]
rows = list(ws.iter_rows(values_only=True))
hdr = {str(v).strip(): i for i, v in enumerate(rows[0]) if v}
c_seller, c_desc, c_thue = hdr["Tên người bán"], hdr["Tên hàng hóa, dịch vụ"], hdr["Tiền thuế"]
c_tt, c_kh, c_so = hdr["Thành tiền chưa thuế"], hdr["Ký hiệu hóa  đơn"], hdr["Số hóa đơn"]

out, ncc_map = [], {}
for i, r in enumerate(rows[1:], start=2):
    seller, desc, thue = r[c_seller], r[c_desc], r[c_thue]
    if seller is None and desc is None:
        continue
    goods = nexia._is_goods(desc)
    if goods:
        sg = {"nguon": "goods"}
    else:
        sg = clf.suggest(seller, desc)
    res = nexia.classify_input_row(clf, kmname, seller, desc, thue)
    s = str(seller or "")
    if s and not CTY.search(s):
        s = ncc_map.setdefault(s, f"NCC-{len(ncc_map) + 1}")
    out.append({
        "i": i, "kyHieu": str(r[c_kh] or "").strip(), "soHd": str(r[c_so] or "").strip(),
        "seller": che(s), "desc": che(desc), "thue": float(thue or 0), "thanhTien": float(r[c_tt] or 0),
        "expected": {"nguon": sg.get("nguon", ""), "kmcp": res["kmcp"], "ten": res["ten"],
                     "tkno": res["tkno"], "tkco": res["tkco"], "vat": res["vat"], "kind": res["kind"]},
    })
json.dump({"rows": out}, open(os.path.join(OUT, "t8-dau-vao.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# Cặp chuẩn hoá — lấy từ chính dữ liệu để phủ đủ dấu/đ/ngoặc/khoảng trắng
samples = sorted({str(x["desc"]) for x in out} | {str(x["seller"]) for x in out})[:300]
samples += ["Đường  Cây Keo", "Vòi sen tắm (Hồng)", "  A   B  ", "Ống  nước Đ/đ", ""]
cases = {
    "norm": [[s, engine.norm(s)] for s in samples],
    "sd": [[s, nexia.sd(s)] for s in samples],
    "hard": [[s, nexia._hard(s)] for s in samples],
}
json.dump(cases, open(os.path.join(OUT, "chuan-hoa-cases.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
from collections import Counter
print("rows:", len(out), "nguon:", Counter(x["expected"]["nguon"] for x in out))

# --- Bổ sung (Task 3, controller quyết): fixture catalog + khoản mục chi phí cho Task 7 ---
# catalog: catalog_map.json của tool (tên, mã, tính chất) — bản chụp catalog_item Masterdata; không PII
cat = [{"ma": c, "ten": t, "tinhChat": tc} for t, c, tc in json.load(open(os.path.join(PKG, "ref", "catalog_map.json"), encoding="utf-8"))]
json.dump(cat, open(os.path.join(OUT, "catalog-t8.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
km = [{"ma": m, "ten": kmname.get(m, ""), "tkNoDefault": tk} for m, tk in json.load(open(os.path.join(PKG, "ref", "km2tk.json"), encoding="utf-8")).items() if not m.startswith("Cần")]
json.dump(km, open(os.path.join(OUT, "kmcp.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("catalog:", len(cat), "kmcp:", len(km))

# --- Đầu ra: expected = classify_output_row (mã nội bộ, độ tin cậy, mã khách). Che: tên người mua cá nhân →
# "KH-<n>" giữ chữ "(shopee)" nếu có; MST → "MST-<n>" theo thứ tự xuất hiện (R12 — dù MST doanh nghiệp không
# phải PII cá nhân, dạng 10 số trùng hình SĐT VN nên scan_pii_secrets.py báo dương tính giả; che bằng pseudo-id
# ổn định thay vì sửa scanner dùng chung). engine coi MST là chuỗi mờ (customerCode = mst.trim()) nên che không
# đổi parity — khớp dim_channel chỉ được kiểm bằng ca dựng tay trong dau-ra.test.ts, không qua golden.
ws_ra = next(ws for ws in openpyxl.load_workbook(XLSX, read_only=True) if nexia.sd("đầu ra") in nexia.sd(ws.title))
# ponytail: workbook T8 không khai <dimension> → ws.max_column = None, nexia._hdr_map() crash;
# tự dựng hmap 1-based từ hàng đầu (cùng công thức _hdr_map) thay vì gọi thẳng hàm đó.
rows_ra_all = list(ws_ra.iter_rows(values_only=True))
hmap = {str(v).strip(): i for i, v in enumerate(rows_ra_all[0], start=1) if v is not None and str(v).strip()}
c_ten, c_mst, c_mua = nexia._find(hmap, "tên hàng"), nexia._find(hmap, "mst người mua"), nexia._find(hmap, "tên người mua")
rows_ra, kh_seq, mst_seq = [], {}, {}
for i, r in enumerate(rows_ra_all[1:], start=2):
    desc, mst, buyer = r[c_ten - 1], r[c_mst - 1], r[c_mua - 1]
    if not desc and not r[nexia._find(hmap, "số hóa đơn") - 1]: continue
    ic, conf, _ = nexia.classify_output_row(desc, mst, buyer)
    b = str(buyer or "")
    b_che = b if CTY.search(b) else kh_seq.setdefault(b.lower(), f"KH-{len(kh_seq)+1}") + (" (shopee)" if "shopee" in b.lower() else "")
    mst_s = str(mst or "").strip()
    mst_che = mst_seq.setdefault(mst_s, f"MST-{len(mst_seq)+1}") if mst_s else ""
    # R13: luật customerCode chốt ở interface Task 7 là mục tiêu, không phải nhãn lịch sử name2kh của Python
    # (name2kh thiếu buyer mới → mặc định 'KHL', không phải hành vi đích) — company (có MST) → mst đã che;
    # cá nhân → 'KHSP' nếu tên chứa "shopee" else 'KHL'.
    makh_exp = mst_che if CTY.search(b) else ("KHSP" if "shopee" in b.lower() else "KHL")
    rows_ra.append({"i": i, "desc": che(desc), "mst": mst_che, "buyer": b_che, "expected": {"ma": ic, "conf": conf, "makh": makh_exp}})
json.dump({"rows": rows_ra}, open(os.path.join(OUT, "t8-dau-ra.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("t8-dau-ra.json:", len(rows_ra), "dòng")
